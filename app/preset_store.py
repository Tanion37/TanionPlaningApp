"""Цепочки Preset: лист presets в tasks.xlsx и появление следующего шага во входящих."""

from __future__ import annotations

import json
import re
from pathlib import Path

from .models import Task
from .tags import INBOX_TAG

SAFE_NAME = re.compile(r"[^\w\-. ()]+", re.UNICODE)


def presets_dir() -> Path:
    from .paths import app_root

    path = app_root() / "data" / "presets"
    path.mkdir(parents=True, exist_ok=True)
    return path


def list_preset_names() -> list[str]:
    names = {path.stem for path in presets_dir().glob("*.json")}
    names.update(_read_preset_sheet().keys())
    return sorted(names, key=str.casefold)


def preset_path(name: str) -> Path:
    stem = SAFE_NAME.sub("", (name or "").strip()).strip(" .") or "preset"
    return presets_dir() / f"{stem}.json"


def ordered_steps(tasks: list[Task]) -> list[Task]:
    """Корень, затем связанные по after_ids. Остальные в конце."""
    by_id = {task.id: task for task in tasks}
    roots = [task for task in tasks if task.prev_id not in by_id]
    if not roots and tasks:
        roots = [tasks[0]]
    seen: set[str] = set()
    ordered: list[Task] = []

    def walk(task: Task) -> None:
        if task.id in seen:
            return
        seen.add(task.id)
        ordered.append(task)
        for child_id in task.after_ids:
            child = by_id.get(child_id)
            if child is not None:
                walk(child)

    for root in roots:
        walk(root)
    for task in tasks:
        if task.id not in seen:
            ordered.append(task)
    return ordered


def steps_from_tasks(tasks: list[Task]) -> list[dict]:
    return [_step_dict(task) for task in ordered_steps(tasks)]


def _step_dict(task: Task) -> dict:
    return {
        "id": task.id,
        "title": (task.title or "").strip(),
        "role": task.role or "",
        "project": task.project or "",
        "description": task.description or "",
        "tags": list(task.tags or []),
        "after": list(task.after_ids),
    }


def save_preset(name: str, tasks: list[Task]) -> str:
    """Записать пресет на лист presets в tasks.xlsx. JSON-файл не создаётся."""
    stem = preset_path(name).stem
    _write_preset_sheet(stem, tasks)
    return stem


PRESET_SHEET = "presets"
_PRESET_HEADERS = ("preset", "id", "название", "роль", "проект", "описание", "теги", "после")


def _xlsx_path() -> Path:
    from .xlsx_store import default_xlsx_path

    return default_xlsx_path()


def _split_cell_list(value) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    parts: list[str] = []
    for chunk in text.replace(";", ",").split(","):
        item = chunk.strip()
        if item and item not in parts:
            parts.append(item)
    return parts


def _sheet_step_rows(name: str, tasks: list[Task]) -> list[list]:
    rows: list[list] = []
    for step in steps_from_tasks(tasks):
        rows.append(
            [
                name,
                str(step.get("id") or ""),
                step.get("title") or "",
                step.get("role") or "",
                step.get("project") or "",
                step.get("description") or "",
                ", ".join(str(tag) for tag in (step.get("tags") or []) if str(tag).strip()),
                ", ".join(str(link) for link in (step.get("after") or []) if str(link).strip()),
            ]
        )
    return rows


def _read_preset_sheet(path: Path | None = None) -> dict[str, list[dict]]:
    """Имя пресета → шаги в порядке строк листа presets."""
    from .xlsx_io import load_workbook_retry

    book = path or _xlsx_path()
    if not book.exists():
        return {}
    try:
        wb = load_workbook_retry(book, data_only=True)
    except (OSError, ValueError):
        return {}
    if PRESET_SHEET not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[PRESET_SHEET]
    grouped: dict[str, list[dict]] = {}
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        cells = list(row or [])
        if index == 0:
            continue
        if not any(cell not in (None, "") for cell in cells):
            continue
        while len(cells) < len(_PRESET_HEADERS):
            cells.append("")
        preset_name = str(cells[0] or "").strip()
        title = str(cells[2] or "").strip()
        if not preset_name or not title:
            continue
        step = {
            "title": title,
            "role": str(cells[3] or ""),
            "project": str(cells[4] or ""),
            "description": str(cells[5] or ""),
            "tags": _split_cell_list(cells[6]),
        }
        step_id = str(cells[1] or "").strip()
        if step_id:
            step["id"] = step_id
        after = _split_cell_list(cells[7])
        if after or step_id:
            step["after"] = after
        grouped.setdefault(preset_name, []).append(step)
    wb.close()
    return grouped


def _replace_preset_rows(name: str, rows: list[list], *, path: Path | None = None) -> None:
    from openpyxl import Workbook

    from .xlsx_io import load_workbook_retry, save_workbook_atomic, xlsx_write_lock

    book = path or _xlsx_path()
    book.parent.mkdir(parents=True, exist_ok=True)
    kept: list[list] = []
    with xlsx_write_lock(book):
        if book.exists():
            wb = load_workbook_retry(book)
            if PRESET_SHEET in wb.sheetnames:
                ws_old = wb[PRESET_SHEET]
                for index, row in enumerate(ws_old.iter_rows(values_only=True)):
                    if index == 0:
                        continue
                    cells = [("" if cell is None else cell) for cell in row]
                    if not cells or str(cells[0] or "").strip() == name:
                        continue
                    kept.append(cells)
                del wb[PRESET_SHEET]
        else:
            wb = Workbook()
            active = wb.active
            active.title = "tasks"
        ws = wb.create_sheet(PRESET_SHEET)
        ws.append(list(_PRESET_HEADERS))
        for cells in kept:
            ws.append(cells)
        for cells in rows:
            ws.append(cells)
        save_workbook_atomic(wb, book)


def _write_preset_sheet(name: str, tasks: list[Task], *, path: Path | None = None) -> None:
    _replace_preset_rows(name, _sheet_step_rows(name, tasks), path=path)


def load_steps(name: str) -> list[dict]:
    stem = Path(name).stem
    sheet_steps = _read_preset_sheet().get(stem)
    if sheet_steps:
        return sheet_steps
    return _steps_from_json(stem)


def _steps_from_json(name: str) -> list[dict]:
    path = presets_dir() / f"{Path(name).stem}.json"
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    raw = data.get("tasks") if isinstance(data, dict) else None
    if not isinstance(raw, list):
        return []
    steps: list[dict] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip()
        if not title:
            continue
        tags = item.get("tags") or []
        step = {
            "title": title,
            "role": str(item.get("role") or ""),
            "project": str(item.get("project") or ""),
            "description": str(item.get("description") or ""),
            "tags": [str(tag) for tag in tags if str(tag).strip()],
        }
        if item.get("id"):
            step["id"] = str(item.get("id"))
        if "after" in item:
            raw_after = item.get("after") or []
            step["after"] = [str(link) for link in raw_after if str(link).strip()]
        steps.append(step)
    return steps


def delete_preset(name: str) -> bool:
    stem = Path(name).stem
    path = presets_dir() / f"{stem}.json"
    removed = False
    if path.exists():
        path.unlink()
        removed = True
    if stem in _read_preset_sheet():
        _replace_preset_rows(stem, [])
        removed = True
    return removed


def rename_preset(old: str, new: str) -> str:
    src = presets_dir() / f"{Path(old).stem}.json"
    dest = preset_path(new)
    if src.exists() and src != dest:
        if dest.exists():
            dest.unlink()
        src.rename(dest)
    old_stem = Path(old).stem
    if old_stem != dest.stem and old_stem in _read_preset_sheet():
        steps = _read_preset_sheet()[old_stem]
        _replace_preset_rows(old_stem, [])
        _write_preset_sheet(dest.stem, tasks_from_steps(steps))
    return dest.stem


def apply_step(task: Task, step: dict, *, preset_name: str, step_index: int) -> None:
    task.title = step["title"]
    task.role = step.get("role") or ""
    task.project = step.get("project") or ""
    task.description = step.get("description") or ""
    task.tags = list(step.get("tags") or [])
    task.preset_name = preset_name
    task.preset_step = step_index


def root_indexes(steps: list[dict]) -> list[int]:
    """Индексы задач без входящей связи. Старый файл без графа — только первая."""
    if not steps:
        return []
    if not any(step.get("id") for step in steps):
        return [0]
    children: set[str] = set()
    for step in steps:
        for link in step.get("after") or []:
            children.add(str(link))
    roots = [index for index, step in enumerate(steps) if str(step.get("id") or "") not in children]
    return roots or [0]


def tasks_from_steps(steps: list[dict]) -> list[Task]:
    """Восстановить граф. Старый файл без id остаётся линейной цепочкой."""
    if not steps:
        return []
    if not any(step.get("id") for step in steps):
        return chain_tasks_from_steps(steps)
    tasks: list[Task] = []
    for index, step in enumerate(steps, start=1):
        task = Task(
            id=str(step.get("id") or f"{index:03d}"),
            title=step["title"],
            role=step.get("role") or "",
            project=step.get("project") or "",
            description=step.get("description") or "",
            tags=list(step.get("tags") or []),
            after_ids=[str(link) for link in (step.get("after") or [])],
        )
        tasks.append(task)
    by_id = {task.id: task for task in tasks}
    for task in tasks:
        task.after_ids = [item for item in task.after_ids if item in by_id and item != task.id]
        task.after_count = len(task.after_ids)
        for child_id in task.after_ids:
            child = by_id[child_id]
            if not child.prev_id:
                child.prev_id = task.id
    return tasks


def chain_tasks_from_steps(steps: list[dict]) -> list[Task]:
    """Линейная цепочка для окна правки: каждая следующая связана с предыдущей."""
    tasks: list[Task] = []
    previous: Task | None = None
    for index, step in enumerate(steps, start=1):
        task = Task(
            id=f"{index:03d}",
            title=step["title"],
            role=step.get("role") or "",
            project=step.get("project") or "",
            description=step.get("description") or "",
            tags=list(step.get("tags") or []),
        )
        if previous is not None:
            task.prev_id = previous.id
            previous.after_ids.append(task.id)
            previous.after_count = len(previous.after_ids)
        tasks.append(task)
        previous = task
    return tasks


def remove_with_outgoing(tasks: list[Task], task_id: str) -> None:
    by_id = {task.id: task for task in tasks}
    drop: set[str] = set()
    stack = [task_id]
    while stack:
        current = stack.pop()
        if current in drop or current not in by_id:
            continue
        drop.add(current)
        stack.extend(by_id[current].after_ids)
    kept: list[Task] = []
    for task in tasks:
        if task.id in drop:
            continue
        task.after_ids = [item for item in task.after_ids if item not in drop]
        task.after_count = len(task.after_ids)
        if task.prev_id in drop:
            task.prev_id = ""
        kept.append(task)
    tasks[:] = kept


def step_index(value) -> int:
    """0 — первый шаг. Пустое значение — шага нет."""
    if value is None or value == "":
        return -1
    try:
        return int(value)
    except (TypeError, ValueError):
        return -1


def _inbox_tags(raw: list) -> list[str]:
    from .tags import CANCEL_TAG, DONE_TAG

    tags = [tag for tag in raw if tag not in {INBOX_TAG, DONE_TAG, CANCEL_TAG}]
    tags.insert(0, INBOX_TAG)
    return tags


def spawn_roots_inbox(store, preset_name: str) -> list:
    """Поставить корневые задачи пресета во входящие. Каждый запуск — новые карточки."""
    name = (preset_name or "").strip()
    steps = load_steps(name)
    created = []
    for index in root_indexes(steps):
        data = steps[index]
        created.append(
            store.add_task(
                data["title"],
                persist=False,
                role=data.get("role") or "",
                project=data.get("project") or "",
                description=data.get("description") or "",
                tags=_inbox_tags(list(data.get("tags") or [])),
                preset_name=name,
                preset_step=index,
                source="preset",
            )
        )
    return created


def _child_indexes(steps: list[dict], step: int) -> list[int]:
    node = steps[step]
    if "after" in node:
        by_id = {str(item.get("id")): index for index, item in enumerate(steps) if item.get("id")}
        found: list[int] = []
        for link in node.get("after") or []:
            index = by_id.get(str(link))
            if index is not None and index not in found:
                found.append(index)
        return found
    nxt = step + 1
    return [nxt] if nxt < len(steps) else []


def _parent_indexes(steps: list[dict], index: int) -> list[int]:
    """Шаги, после которых стоит этот. Без графа — предыдущий шаг линейной цепочки."""
    if not any("after" in step for step in steps):
        return [index - 1] if index > 0 else []
    node_id = str(steps[index].get("id") or "")
    if not node_id:
        return []
    parents: list[int] = []
    for parent_index, step in enumerate(steps):
        links = [str(link) for link in (step.get("after") or [])]
        if node_id in links and parent_index not in parents:
            parents.append(parent_index)
    return parents


def _same_run(store, task: Task) -> list[Task]:
    """Живые задачи одного запуска пресета: общий предок или связь prev/after."""
    name = (getattr(task, "preset_name", "") or "").strip()
    pool = [item for item in store.tasks if (item.preset_name or "") == name]
    by_id = {item.id: item for item in pool}
    seen: set[str] = set()
    current: Task | None = task if task.id in by_id else None
    guard = 0
    while current is not None and current.id not in seen and guard < 100:
        seen.add(current.id)
        current = by_id.get((current.prev_id or "").strip())
        guard += 1
    changed = True
    while changed:
        changed = False
        for item in pool:
            if item.id in seen:
                for link in item.after_ids:
                    if link in by_id and link not in seen:
                        seen.add(link)
                        changed = True
                continue
            if (item.prev_id or "") in seen:
                seen.add(item.id)
                changed = True
    return [by_id[item_id] for item_id in seen if item_id in by_id]


def spawn_linked_inbox(store, task: Task) -> list:
    """Следующие задачи во входящих. Общий потомок ждёт, пока выполнены все предки."""
    name = (getattr(task, "preset_name", "") or "").strip()
    step = step_index(getattr(task, "preset_step", -1))
    if not name or step < 0:
        return []
    steps = load_steps(name)
    if step >= len(steps):
        return []
    run = _same_run(store, task)
    live_steps: dict[int, list[Task]] = {}
    for item in run:
        live_steps.setdefault(step_index(item.preset_step), []).append(item)
    created = []
    for index in _child_indexes(steps, step):
        if any(step_index(item.preset_step) == index for item in run):
            continue
        parents = _parent_indexes(steps, index)
        if not parents:
            continue
        ready = True
        parent_live: list[Task] = []
        for parent_index in parents:
            lives = live_steps.get(parent_index) or []
            if not lives or not all(item.is_done() for item in lives):
                ready = False
                break
            parent_live.extend(lives)
        if not ready:
            continue
        data = steps[index]
        child = store.add_task(
            data["title"],
            persist=False,
            role=data.get("role") or "",
            project=data.get("project") or "",
            description=data.get("description") or "",
            tags=_inbox_tags(list(data.get("tags") or [])),
            preset_name=name,
            preset_step=index,
            prev_id=task.id,
            source="preset",
        )
        for parent in parent_live:
            if child.id not in parent.after_ids:
                parent.after_ids.append(child.id)
                parent.after_count = len(parent.after_ids)
        created.append(child)
        run.append(child)
    return created


def spawn_next_inbox(store, task: Task):
    """Совместимость: первая из следующих связанных задач."""
    created = spawn_linked_inbox(store, task)
    return created[0] if created else None
