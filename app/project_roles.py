"""Проект хранит, какой исполнитель занимает роль. Preset задаёт только сами роли."""

from __future__ import annotations

from pathlib import Path

from .models import Task
from .projects import preset_of_project, project_key, same_project
from .roles import StudioRole
from .tags import INBOX_TAG

PROJECT_LEAD = StudioRole.PROJECT_LEAD.value
ASSIGN_PREFIX = "Назначить роль «"
ASSIGN_BOTH = "Назначить роль и исполнителя"
SHEET_NAME = "project_roles"
_HEADERS = ("проект", "роль", "исполнитель")


def assignment_title(role: str) -> str:
    return f"{ASSIGN_PREFIX}{(role or '').strip()}»"


def role_from_assignment(task: Task) -> str:
    title = (getattr(task, "title", "") or "").strip()
    if not title.startswith(ASSIGN_PREFIX) or not title.endswith("»"):
        return ""
    return title[len(ASSIGN_PREFIX) : -1].strip()


def is_role_and_executor_task(task: Task) -> bool:
    return (getattr(task, "title", "") or "").strip() == ASSIGN_BOTH


_ROLE_KEYS = ("Роль", "роль")
_EXECUTOR_KEYS = ("Исполнитель", "исполнитель")


def _meta_any(task: Task, keys: tuple[str, ...]) -> str:
    for line in (getattr(task, "description", "") or "").splitlines():
        for key in keys:
            prefix = f"{key}:"
            if line.startswith(prefix):
                return line[len(prefix) :].strip()
    return ""


def _is_choice_line(line: str) -> bool:
    for key in (*_ROLE_KEYS, *_EXECUTOR_KEYS):
        if line.startswith(f"{key}:"):
            return True
    return False


def strip_choice_lines(text: str) -> str:
    lines = [line for line in (text or "").splitlines() if not _is_choice_line(line)]
    return "\n".join(lines).strip()


def description_with_choices(text: str, role: str, person: str) -> str:
    """Текст задачи и строки «Роль: …» / «Исполнитель: …» в конце."""
    base = strip_choice_lines(text)
    choice: list[str] = []
    role = (role or "").strip()
    person = (person or "").strip()
    if role:
        choice.append(f"Роль: {role}")
    if person:
        choice.append(f"Исполнитель: {person}")
    parts = [part for part in (base, "\n".join(choice)) if part]
    return "\n\n".join(parts).strip()


def description_without_choices(task: Task) -> str:
    """Текст задачи без строк «Роль:» и «Исполнитель:»."""
    return strip_choice_lines(getattr(task, "description", "") or "")


def _write_choices(task: Task, role: str, person: str) -> None:
    task.description = description_with_choices(getattr(task, "description", "") or "", role, person)


def set_description_keep_choices(
    task: Task,
    text: str,
    role: str | None = None,
    person: str | None = None,
) -> None:
    if role is None:
        role = pending_role(task)
    if person is None:
        person = pending_executor(task)
    task.description = text or ""
    _write_choices(task, role, person)


def pending_role(task: Task) -> str:
    return _meta_any(task, _ROLE_KEYS)


def pending_executor(task: Task) -> str:
    return _meta_any(task, _EXECUTOR_KEYS)


def _choices_are_canonical(task: Task, role: str, person: str) -> bool:
    current = getattr(task, "description", "") or ""
    return current == description_with_choices(current, role, person)


def remember_role(task: Task, role: str) -> bool:
    role = (role or "").strip()
    person = pending_executor(task)
    if pending_role(task) == role and _choices_are_canonical(task, role, person):
        return False
    _write_choices(task, role, person)
    return True


def remember_executor(task: Task, executor: str) -> bool:
    executor = (executor or "").strip()
    role = pending_role(task)
    if pending_executor(task) == executor and _choices_are_canonical(task, role, executor):
        return False
    _write_choices(task, role, executor)
    return True


def pin_assignment_lead(store, task: Task) -> None:
    """Карточку «Назначить роль и исполнителя» держит руководитель проекта."""
    if not is_role_and_executor_task(task):
        return
    task.role = PROJECT_LEAD
    lead = executor_for(getattr(task, "project", ""), PROJECT_LEAD, getattr(store, "path", None))
    if lead:
        task.executor = lead


def hold_executor_without_role(store, task: Task, person: str) -> bool | None:
    """Проект не запрещает назначить исполнителя.

    Раньше задача с проектом и без роли не получала чужого исполнителя:
    выбор уходил на карточку «Назначить роль и исполнителя», а поле исполнителя
    очищалось. Из-за этого задачу и проект нельзя было передать, в том числе
    человеку, у которого уже есть другой проект. None — назначение делает вызывающий код.
    """
    del store, task, person
    return None


def _finish_assignment(store, task: Task, role: str, person: str, path: Path | None) -> bool:
    from datetime import date

    from .tags import ACTUAL_TAG, CANCEL_TAG, CONTROL_TAG, DONE_TAG, INBOX_TAG

    if task.is_done() and task.has_tag(DONE_TAG):
        return False
    source = store.get(task.prev_id) if task.prev_id and hasattr(store, "get") else None
    if source is not None and not getattr(source, "deleted", 0) and source.id != task.id:
        if not is_role_and_executor_task(source) and not role_from_assignment(source):
            source.role = role
            from .day_tasks import apply_executor_assignment

            apply_executor_assignment(source, person)
            if hasattr(store, "mark_changed"):
                store.mark_changed(source)
    project = (task.project or (getattr(source, "project", "") if source else "") or "").strip()
    if project:
        assign_executor(project, role, person, path)
        apply_executor_to_tasks(store, project, role, person)
    task.remove_tag(CANCEL_TAG)
    task.remove_tag(INBOX_TAG)
    task.remove_tag(ACTUAL_TAG)
    task.remove_tag(CONTROL_TAG)
    task.add_tag(DONE_TAG)
    task.completed_at = date.today()
    if hasattr(store, "mark_changed"):
        store.mark_changed(task)
    return True


def settle_assignment(store, task: Task, path: Path | None = None) -> bool:
    """Когда выбор полный — записать его и сдать карточку в выполненные."""
    if is_role_and_executor_task(task):
        pin_assignment_lead(store, task)
        role = pending_role(task)
        person = pending_executor(task)
    else:
        role = role_from_assignment(task)
        person = pending_executor(task)
    if not role or not person:
        return False
    return _finish_assignment(store, task, role, person, path)


def _path(path: Path | None) -> Path:
    if path is not None:
        return path
    from .xlsx_store import default_xlsx_path

    return default_xlsx_path()


def load_assignments(path: Path | None = None) -> dict[str, dict[str, str]]:
    """Имя проекта → {роль: исполнитель}."""
    from .xlsx_io import load_workbook_retry

    book = _path(path)
    if not book.exists():
        return {}
    try:
        wb = load_workbook_retry(book, data_only=True)
    except (OSError, ValueError):
        return {}
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        return {}
    found: dict[str, dict[str, str]] = {}
    ws = wb[SHEET_NAME]
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if index == 0:
            continue
        cells = list(row or [])
        project = str(cells[0] or "").strip() if cells else ""
        role = str(cells[1] or "").strip() if len(cells) > 1 else ""
        executor = str(cells[2] or "").strip() if len(cells) > 2 else ""
        if not project or not role or not executor:
            continue
        bucket = found.setdefault(project, {})
        bucket[role] = executor
    wb.close()
    return found


def _bucket(project: str, path: Path | None = None) -> tuple[str, dict[str, str]]:
    mapping = load_assignments(path)
    for name, roles in mapping.items():
        if same_project(name, project):
            return name, dict(roles)
    return (project or "").strip(), {}


def executor_for(project: str, role: str, path: Path | None = None) -> str:
    _name, roles = _bucket(project, path)
    return roles.get((role or "").strip(), "")


def roles_of_project(project: str, tasks: list, path: Path | None = None) -> list[str]:
    """Роли образца и уже назначенные роли этого проекта. Руководитель всегда первый."""
    ordered: list[str] = []

    def add(role: str) -> None:
        text = (role or "").strip()
        if text and text not in ordered:
            ordered.append(text)

    add(PROJECT_LEAD)
    preset = preset_of_project(tasks, project)
    if preset:
        from .preset_store import load_steps

        for step in load_steps(preset):
            add(str(step.get("role") or ""))
    for task in tasks or []:
        if same_project(getattr(task, "project", ""), project):
            add(getattr(task, "role", ""))
    _name, assigned = _bucket(project, path)
    for role in assigned:
        add(role)
    return ordered


def _write(rows: list[tuple[str, str, str]], path: Path | None = None) -> None:
    from openpyxl import Workbook

    from .xlsx_io import load_workbook_retry, save_workbook_atomic, xlsx_write_lock

    book = _path(path)
    book.parent.mkdir(parents=True, exist_ok=True)
    with xlsx_write_lock(book):
        if book.exists():
            wb = load_workbook_retry(book)
            if SHEET_NAME in wb.sheetnames:
                del wb[SHEET_NAME]
        else:
            wb = Workbook()
            active = wb.active
            active.title = "tasks"
        ws = wb.create_sheet(SHEET_NAME)
        ws.append(list(_HEADERS))
        for project, role, executor in rows:
            ws.append([project, role, executor])
        save_workbook_atomic(wb, book)


def _rows(path: Path | None = None) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    for project, roles in load_assignments(path).items():
        for role, executor in roles.items():
            rows.append((project, role, executor))
    return rows


def assign_executor(project: str, role: str, executor: str, path: Path | None = None) -> None:
    project = (project or "").strip()
    role = (role or "").strip()
    executor = (executor or "").strip()
    if not project or not role or not executor:
        return
    kept = [
        (name, item_role, person)
        for name, item_role, person in _rows(path)
        if not (same_project(name, project) and item_role == role)
    ]
    display = project
    for name, _role, _person in kept:
        if same_project(name, project):
            display = name
            break
    kept.append((display, role, executor))
    _write(kept, path)


def rename_project(old: str, new: str, path: Path | None = None) -> None:
    old = (old or "").strip()
    new = (new or "").strip()
    if not old or not new or project_key(old) == project_key(new):
        return
    rows = [
        (new if same_project(name, old) else name, role, person)
        for name, role, person in _rows(path)
    ]
    _write(rows, path)


def forget_project(project: str, path: Path | None = None) -> None:
    rows = [
        (name, role, person)
        for name, role, person in _rows(path)
        if not same_project(name, project)
    ]
    _write(rows, path)


def apply_executor_to_tasks(store, project: str, role: str, executor: str) -> None:
    """Задачам этой роли в проекте поставить исполнителя, даже если он уже был другим.

    Руководитель проекта забирает и задачи проекта без роли.
    Другой проект у этого человека назначению не мешает.
    """
    from .day_tasks import apply_executor_assignment

    person = (executor or "").strip()
    role = (role or "").strip()
    if not person or not role:
        return
    for task in store.tasks:
        if getattr(task, "deleted", 0):
            continue
        if not same_project(getattr(task, "project", ""), project):
            continue
        task_role = (getattr(task, "role", "") or "").strip()
        if task_role != role and not (role == PROJECT_LEAD and not task_role):
            continue
        if is_role_and_executor_task(task) or role_from_assignment(task):
            continue
        if (getattr(task, "executor", "") or "").strip() == person:
            continue
        apply_executor_assignment(task, person, store=store)
        if hasattr(store, "mark_changed"):
            store.mark_changed(task)


def _open_assignment(store, project: str, role: str) -> Task | None:
    title = assignment_title(role)
    for task in store.tasks:
        if getattr(task, "deleted", 0) or task.is_done() or task.is_cancelled():
            continue
        if (task.title or "").strip() != title:
            continue
        if same_project(task.project, project):
            return task
    return None


def _open_both_assignment(store, source: Task) -> Task | None:
    for task in store.tasks:
        if getattr(task, "deleted", 0) or task.is_done() or task.is_cancelled():
            continue
        if not is_role_and_executor_task(task):
            continue
        if (task.prev_id or "") == source.id:
            return task
    return None


def bind_task_executor(store, task: Task, path: Path | None = None) -> list[Task]:
    """Поставить исполнителя роли. Если роли нет — руководитель назначает роль и исполнителя."""
    if is_role_and_executor_task(task) or role_from_assignment(task):
        return []
    role = (task.role or "").strip()
    project = (task.project or "").strip()
    if not project:
        return []
    if not role:
        lead = executor_for(project, PROJECT_LEAD, path)
        if not lead or _open_both_assignment(store, task) is not None:
            return []
        created = store.add_task(
            ASSIGN_BOTH,
            persist=False,
            executor=lead,
            role=PROJECT_LEAD,
            project=project,
            description=f"У задачи «{task.title}» в проекте «{project}» не выбраны роль и исполнитель.",
            tags=[INBOX_TAG],
            prev_id=task.id,
            source="app",
        )
        return [created]
    person = executor_for(project, role, path)
    if person:
        if not (task.executor or "").strip():
            task.executor = person
        return []
    already = (task.executor or "").strip()
    if already and role != PROJECT_LEAD:
        assign_executor(project, role, already, path)
        return []
    if role == PROJECT_LEAD:
        return []
    lead = executor_for(project, PROJECT_LEAD, path)
    if not lead or _open_assignment(store, project, role) is not None:
        return []
    created = store.add_task(
        assignment_title(role),
        persist=False,
        executor=lead,
        role=PROJECT_LEAD,
        project=project,
        description=f"В проекте «{project}» назначьте исполнителя на роль «{role}».",
        tags=[INBOX_TAG],
        source="app",
    )
    return [created]
