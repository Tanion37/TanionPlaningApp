"""Атомарная запись xlsx и межпроцессный замок (планировщик + бот)."""

from __future__ import annotations

import os
import sys
import tempfile
import time
from contextlib import contextmanager
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

LOCK_WAIT_SEC = 15.0
LOAD_ATTEMPTS = 6
LOAD_RETRY_SEC = 0.12


def _is_transient_xlsx_error(exc: BaseException) -> bool:
    if isinstance(exc, BadZipFile):
        return True
    if isinstance(exc, PermissionError):
        return True
    msg = str(exc).casefold()
    return "magic" in msg or "zip file" in msg or "permission" in msg


def load_workbook_retry(path: Path, *, data_only: bool = False):
    """Открыть xlsx; при битой шапке/блокировке – короткие повторы."""
    last: BaseException | None = None
    for attempt in range(LOAD_ATTEMPTS):
        try:
            return load_workbook(path, data_only=data_only)
        except Exception as exc:  # noqa: BLE001 — отличить гонку записи от прочих ошибок
            last = exc
            if not _is_transient_xlsx_error(exc) or attempt + 1 >= LOAD_ATTEMPTS:
                raise
            time.sleep(LOAD_RETRY_SEC * (attempt + 1))
    assert last is not None
    raise last


def save_workbook_atomic(wb, path: Path) -> None:
    """Писать во временный файл и подменить – читатель не видит обрезанный zip."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=".tmp-xlsx-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except Exception:
        tmp.unlink(missing_ok=True)
        raise


@contextmanager
def xlsx_write_lock(path: Path, *, wait_sec: float = LOCK_WAIT_SEC):
    """Один писатель на файл: load+правка+save под замком."""
    path = Path(path)
    lock_path = path.with_name(path.name + ".lock")
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    deadline = time.time() + wait_sec
    with open(lock_path, "a+b") as handle:
        if sys.platform == "win32":
            import msvcrt

            while True:
                handle.seek(0)
                try:
                    msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError:
                    if time.time() >= deadline:
                        raise TimeoutError(f"Не удалось взять замок {lock_path}") from None
                    time.sleep(0.05)
            try:
                yield
            finally:
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl

            while True:
                try:
                    fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                    break
                except OSError:
                    if time.time() >= deadline:
                        raise TimeoutError(f"Не удалось взять замок {lock_path}") from None
                    time.sleep(0.05)
            try:
                yield
            finally:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
